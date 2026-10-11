#!/usr/bin/env python3
"""Audit a .pptx deck against the slide-design-system tokens.

Usage:
    python audit_deck.py deck.pptx [--tokens ../assets/tokens.json]

Reports fonts, sizes, and colors that are not in the token set, text below the minimum
size, and per-slide title consistency. Read-only: it never modifies the deck.
Requires: pip install python-pptx
"""
import argparse
import json
import os
import sys
from collections import Counter

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor


def load_tokens(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def iter_shapes(shapes):
    for shp in shapes:
        yield shp
        if shp.shape_type == 6 and hasattr(shp, "shapes"):  # group
            yield from iter_shapes(shp.shapes)


def run_color(run):
    try:
        if run.font.color and run.font.color.type is not None and run.font.color.rgb is not None:
            return "#" + str(run.font.color.rgb).upper()
    except (AttributeError, TypeError):
        pass
    return None


def fill_color(shape):
    try:
        if shape.fill.type == 1 and shape.fill.fore_color.rgb is not None:
            return "#" + str(shape.fill.fore_color.rgb).upper()
    except (AttributeError, TypeError):
        pass
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    default_tokens = os.path.join(os.path.dirname(__file__), "..", "assets", "tokens.json")
    ap.add_argument("--tokens", default=default_tokens)
    args = ap.parse_args()

    tokens = load_tokens(args.tokens)
    allowed_sizes = {t["size"] for t in tokens["type"].values()}
    allowed_colors = {c.upper() for c in tokens["color"].values()}
    allowed_colors |= {c.upper() for c in tokens.get("terminal", {}).values() if isinstance(c, str)}
    allowed_colors |= {c.upper() for c in tokens.get("terminal", {}).get("dots", [])}
    allowed_fonts = {f.lower() for f in tokens["font"]["text"] + tokens["font"]["code"] + tokens["font"].get("wordmark", [])}
    min_body = tokens["limits"]["min_body_pt"]

    prs = Presentation(args.deck)
    fonts, sizes, colors, fills = Counter(), Counter(), Counter(), Counter()
    findings = []
    title_positions = []

    for idx, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        if title_shape is not None:
            title_positions.append((idx, title_shape.left, title_shape.top, title_shape.width))
        elif idx > 1:
            findings.append(f"Slide {idx}: no title placeholder")

        if not (slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None
                and slide.notes_slide.notes_text_frame.text.strip()):
            findings.append(f"Slide {idx}: no speaker notes")

        for shp in iter_shapes(slide.shapes):
            fc = fill_color(shp)
            if fc:
                fills[fc] += 1
                if fc not in allowed_colors:
                    findings.append(f"Slide {idx}: off-token fill {fc} on '{shp.name}'")
            if not shp.has_text_frame:
                continue
            for para in shp.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    name = run.font.name
                    if name:
                        fonts[name] += 1
                        if name.lower() not in allowed_fonts:
                            findings.append(f"Slide {idx}: font '{name}' not in token stacks")
                    if run.font.size:
                        pt = round(run.font.size.pt, 1)
                        sizes[pt] += 1
                        if pt not in allowed_sizes:
                            findings.append(f"Slide {idx}: size {pt} pt is not a token size")
                        if pt < 12:
                            findings.append(f"Slide {idx}: text below 12 pt ({pt})")
                    col = run_color(run)
                    if col:
                        colors[col] += 1
                        if col not in allowed_colors:
                            findings.append(f"Slide {idx}: off-token text color {col}")

    print(f"Slides: {len(prs.slides)}")
    print(f"Fonts used: {dict(fonts) or 'theme default only'}")
    print(f"Distinct font families: {len(fonts)} (target: at most 2, one text and one code)")
    print(f"Sizes used (pt): {dict(sorted(sizes.items()))}")
    print(f"Distinct sizes: {len(sizes)} (target: at most {len(allowed_sizes)})")
    print(f"Text colors: {dict(colors)}")
    print(f"Fill colors: {dict(fills)}")

    if len(title_positions) > 1:
        tops = {t[2] for t in title_positions}
        lefts = {t[1] for t in title_positions}
        if len(tops) > 1 or len(lefts) > 1:
            findings.append(
                "Title placeholders differ in position across slides "
                f"({len(tops)} distinct tops, {len(lefts)} distinct lefts)"
            )

    # Deduplicate findings while preserving order
    seen, unique = set(), []
    for f in findings:
        if f not in seen:
            seen.add(f)
            unique.append(f)

    print()
    if unique:
        print(f"{len(unique)} finding(s):")
        for f in unique:
            print(f"  - {f}")
        sys.exit(1)
    print("No findings. Deck matches the token set.")


if __name__ == "__main__":
    main()
